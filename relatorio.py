"""
Geração dos arquivos de saída:
  - CSV idêntico ao RELATORIO-FINAL-UNIFICADO.csv de hoje (';' e vírgula decimal)
  - Excel (.xlsx) com valores numéricos reais, formatação R$ e cabeçalho com os parâmetros
"""
import io

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


COLUNAS_VALORES = [
    "valor_principal_proposta", "juros_proposta", "Total_data_pagto",
    "Valor_Depositado", "Per_Prin_Pagto", "Per_Juros_Pagto",
    "Saldo_Principal", "Saldo_Juros", "Saldo_Total",
    "juros_selic_proposta", "Per_Selic_Pagto", "Saldo_Selic",
]
COLUNAS_INDICES = ["IPCA_Selic_Aplicado", "Juros_Aplicado"]


def formatar_indice(val):
    if isinstance(val, (int, float)):
        if val == 0:
            return "0"
        s = f"{val:.9f}".rstrip("0")
        if s.endswith("."):
            s = s[:-1]
        return s.replace(".", ",")
    return val


def formatar_moeda(x):
    return f"R$ {x:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def gerar_csv(df_final: pd.DataFrame) -> bytes:
    df_csv = df_final.copy()
    for c in COLUNAS_INDICES:
        df_csv[c] = df_csv[c].apply(formatar_indice)
    buf = io.StringIO()
    df_csv.to_csv(buf, sep=";", decimal=",", index=False)
    return buf.getvalue().encode("utf-8-sig")


def tabela_para_tela(df_final: pd.DataFrame) -> pd.DataFrame:
    """Mesma formatação visual que o script imprimia no terminal."""
    df = df_final.copy()
    for c in [c for c in COLUNAS_VALORES if c in df.columns]:
        df[c] = df[c].apply(formatar_moeda)
    for c in COLUNAS_INDICES:
        df[c] = df[c].apply(formatar_indice)
    df["ordem_pagamento"] = df["ordem_pagamento"].astype(str)
    return df


def gerar_excel(df_final: pd.DataFrame, parametros: dict, pagamentos: pd.DataFrame) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Relatorio"

    azul = PatternFill("solid", fgColor="1F3864")
    cinza = PatternFill("solid", fgColor="F2F2F2")
    fina = Side(style="thin", color="BFBFBF")
    borda = Border(left=fina, right=fina, top=fina, bottom=fina)
    colunas = list(df_final.columns)
    n_cols = len(colunas)

    ws.cell(row=1, column=1, value="RELATÓRIO HISTÓRICO UNIFICADO - ESTEIRA DE ATUALIZAÇÃO").font = Font(bold=True, size=14)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n_cols)

    linha = 3
    for rotulo, valor in parametros.items():
        ws.cell(row=linha, column=1, value=rotulo).font = Font(bold=True)
        c = ws.cell(row=linha, column=3, value=valor)
        if isinstance(valor, float):
            c.number_format = '"R$" #,##0.00'
        linha += 1

    linha += 1
    cab = linha
    for j, nome in enumerate(colunas, start=1):
        c = ws.cell(row=cab, column=j, value=nome)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = azul
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = borda
    ws.row_dimensions[cab].height = 32

    for i, (_, r) in enumerate(df_final.iterrows(), start=1):
        lin = cab + i
        for j, nome in enumerate(colunas, start=1):
            v = r[nome]
            if hasattr(v, "item"):
                v = v.item()
            c = ws.cell(row=lin, column=j, value=v)
            c.border = borda
            if nome in COLUNAS_VALORES:
                c.number_format = '"R$" #,##0.00'
            elif nome in COLUNAS_INDICES:
                c.number_format = "0.000000000"
            if r["ordem_pagamento"] == "FINAL":
                c.font = Font(bold=True)
                c.fill = cinza

    larguras = {"ordem_pagamento": 10, "data_pagamento": 10, "Regra_Ativa": 18,
                "IPCA_Selic_Aplicado": 14, "Juros_Aplicado": 12}
    for j, w in enumerate((larguras.get(c, 16) for c in colunas), start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = ws.cell(row=cab + 1, column=1)

    # Aba com os dados de entrada (para conferência/auditoria)
    ws2 = wb.create_sheet("Entrada")
    ws2.append(["Ordem_Pagamento", "Data de Pagamento", "Valor Depositado"])
    for cel in ws2[1]:
        cel.font = Font(bold=True)
    for _, p in pagamentos.iterrows():
        ws2.append([int(p["ordem_pagamento"]), p["data_pagamento"].strftime("%d/%m/%Y"),
                    float(p["Valor_Depositado"])])
        ws2.cell(row=ws2.max_row, column=3).number_format = '"R$" #,##0.00'
    for col, w in zip("ABC", [16, 18, 18]):
        ws2.column_dimensions[col].width = w

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
