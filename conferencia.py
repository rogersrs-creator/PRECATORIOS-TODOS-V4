"""
Conferência automática dos motores com casos homologados.

Cada caso é um relatório Excel no formato de saída do próprio app (abas "Relatorio" e "Entrada"),
guardado na pasta conferencia/ do repositório. Para incluir um caso novo, basta colocar o arquivo
nessa pasta.

Para cada caso, o app refaz o cálculo com os índices e os motores atuais e compara TODAS as linhas
(total na data de cada pagamento e saldo após cada pagamento) com o relatório homologado.
"""
import glob
import io
import os

import openpyxl
import pandas as pd

PASTA = os.path.join(os.path.dirname(os.path.abspath(__file__)), "conferencia")
TOLERANCIA = 0.05     # R$ - diferenças de arredondamento aceitas
COLUNAS_COMPARADAS = ["Total_data_pagto", "Saldo_Total"]


def ler_relatorio(origem, nome=None):
    """Lê um relatório do app (caminho ou bytes). Retorna o caso com entradas e linhas esperadas."""
    arquivo = io.BytesIO(origem) if isinstance(origem, (bytes, bytearray)) else origem
    wb = openpyxl.load_workbook(arquivo, data_only=True)
    ws = wb["Relatorio"]
    rows = [[c.value for c in r] for r in ws.iter_rows()]
    hdr = next(i for i, r in enumerate(rows) if r[0] == "ordem_pagamento")
    params = {r[0]: r[2] for r in rows[1:hdr] if r[0] and r[2] is not None}
    cols = [c for c in rows[hdr] if c]
    esperado = pd.DataFrame([r[:len(cols)] for r in rows[hdr + 1:] if r[0] is not None], columns=cols)
    ent = [r for r in wb["Entrada"].iter_rows(min_row=2, values_only=True) if r[0] is not None] \
        if "Entrada" in wb.sheetnames else []
    pagamentos = pd.DataFrame({
        "ordem_pagamento": [int(r[0]) for r in ent],
        "data_pagamento": pd.to_datetime([r[1] for r in ent], dayfirst=True),
        "Valor_Depositado": [float(r[2]) for r in ent],
    }, columns=["ordem_pagamento", "data_pagamento", "Valor_Depositado"])
    return {
        "nome": nome or os.path.basename(str(origem)),
        "esfera": params["Esfera"],
        "natureza": params.get("Natureza", "Não Tributário"),
        "data_base": pd.to_datetime(params["Data-base inicial"], dayfirst=True).replace(day=1),
        "ano": int(params["Ano da proposta"]),
        "principal": float(params["Valor principal inicial"]),
        "juros": float(params.get("Juros iniciais", 0) or 0),
        "selic": float(params.get("Juros Selic iniciais", 0) or 0),
        "data_final": pd.to_datetime(params["Data final da dívida"], dayfirst=True).replace(day=1),
        "pagamentos": pagamentos,
        "esperado": esperado,
    }


def listar_casos():
    return sorted(glob.glob(os.path.join(PASTA, "*.xlsx")))


def conferir(caso, bases, df_graca, regra_para, processar_esteira):
    """Refaz o cálculo de um caso e compara linha a linha. Retorna (resumo, detalhe)."""
    resumo = {"Caso": caso["nome"], "Esfera": caso["esfera"], "Natureza": caso.get("natureza", "Não Tributário"),
              "Proposta": caso["ano"],
              "Pagamentos": len(caso["pagamentos"])}
    try:
        regra = regra_para(caso["esfera"], caso["ano"], caso.get("natureza", "Não Tributário"))
        linha = df_graca[df_graca["ANO_PROPOSTA"] == caso["ano"]]
        if regra is None or linha.empty:
            raise ValueError("sem regime ou sem período de graça para esta proposta")
        linha = linha.iloc[0]
        # Não tributário: data-base da planilha GRACA (confere também a planilha). Tributário: a do próprio caso.
        data_base = caso["data_base"] if getattr(regra, "data_base_propria", False) \
            else linha["DATA_BASE_INICIAL_CALCULOS"]
        calc = processar_esteira(regra, caso["pagamentos"], caso["principal"], caso["juros"],
                                 data_base, linha["FIM_PER_GRACA"],
                                 caso["data_final"], bases, selic_inicial=caso["selic"])
    except Exception as e:  # noqa: BLE001
        resumo.update({"Homologado": None, "App": None, "Maior diferença": None,
                       "Situação": f"❌ erro: {e}"})
        return resumo, None

    esp = caso["esperado"]
    n = min(len(calc), len(esp))
    detalhe = pd.DataFrame({
        "Linha": calc["ordem_pagamento"].astype(str).values[:n],
        "Mês": calc["data_pagamento"].values[:n],
        "Total na data – homologado": esp["Total_data_pagto"].astype(float).values[:n],
        "Total na data – app": calc["Total_data_pagto"].astype(float).values[:n],
        "Saldo – homologado": esp["Saldo_Total"].astype(float).values[:n],
        "Saldo – app": calc["Saldo_Total"].astype(float).values[:n],
    })
    detalhe["Diferença"] = (detalhe[["Total na data – app", "Saldo – app"]].values
                            - detalhe[["Total na data – homologado", "Saldo – homologado"]].values
                            ).__abs__().max(axis=1)
    maior = float(detalhe["Diferença"].max()) if n else 0.0
    linhas_ok = len(calc) == len(esp)
    ok = linhas_ok and maior <= TOLERANCIA
    if not linhas_ok:
        situacao = f"❌ nº de linhas diferente ({len(esp)} homologadas × {len(calc)} no app)"
    elif ok:
        situacao = "✅ confere"
    else:
        primeira = detalhe[detalhe["Diferença"] > TOLERANCIA].iloc[0]
        situacao = f"❌ diverge a partir da linha {primeira['Linha']} ({primeira['Mês']})"
    resumo.update({"Homologado": float(esp["Saldo_Total"].iloc[-1]), "App": float(calc["Saldo_Total"].iloc[-1]),
                   "Maior diferença": maior, "Situação": situacao})
    return resumo, detalhe


def rodar_todos(casos, bases, df_graca, regra_para, processar_esteira):
    resumos, detalhes = [], {}
    for caso in casos:
        r, d = conferir(caso, bases, df_graca, regra_para, processar_esteira)
        resumos.append(r)
        detalhes[caso["nome"]] = d
    return pd.DataFrame(resumos), detalhes


# ------------------------------------------------------------------------------
# Verificação interna das bases de índices (independe dos casos)
# ------------------------------------------------------------------------------
# Os motores usam o número-índice só nas pontas de cada período; um erro num mês intermediário só
# apareceria num caso que usasse aquele mês. Por isso conferimos todas as linhas: cada número-índice
# precisa bater com o anterior e com a variação do mês informada na mesma linha.
REGRAS_BASES = {
    "IPCA-15":  ("ipca15", "mult"),      # n = n_anterior × (1 + variação/100)
    "Poupança": ("poup", "soma_pct"),    # n = n_anterior + variação (em %)
    "SELIC":    ("selic", "soma_dec"),   # n = n_anterior + variação/100
    "IPCA":     ("ipca", "mult"),
}
TOL_BASE = 1e-6


def verificar_bases(bases):
    """Retorna (resumo por base, lista de problemas encontrados)."""
    resumo, problemas = [], []
    for nome, (campo, tipo) in REGRAS_BASES.items():
        df = getattr(bases, campo).sort_values("data").reset_index(drop=True)
        n_prob = 0
        # meses repetidos ou faltando
        dup = df[df["data"].duplicated()]
        for d in dup["data"]:
            problemas.append({"Base": nome, "Mês": d.strftime("%m/%Y"), "Problema": "mês repetido"}); n_prob += 1
        esperado = pd.date_range(df["data"].min(), df["data"].max(), freq="MS")
        for d in sorted(set(esperado) - set(df["data"])):
            problemas.append({"Base": nome, "Mês": d.strftime("%m/%Y"), "Problema": "mês faltando"}); n_prob += 1
        if "variacao" in df.columns:
            n, v, ant = df["n_indice"], df["variacao"], df["n_indice"].shift(1)
            if tipo == "mult":
                calc = ant * (1 + v / 100); erro = (n - calc).abs() / n.abs().where(n != 0, 1)
            elif tipo == "soma_pct":
                calc = ant + v; erro = (n - calc).abs()
            else:
                calc = ant + v / 100; erro = (n - calc).abs()
            # variação que o próprio número-índice indica (na mesma unidade da coluna de variação)
            if tipo == "mult":
                implicita = (n / ant - 1) * 100
            elif tipo == "soma_pct":
                implicita = n - ant
            else:
                implicita = (n - ant) * 100
            for i in erro.index[1:]:
                if pd.isna(erro[i]) or erro[i] > TOL_BASE:
                    fmt = lambda x: "vazia" if pd.isna(x) else f"{x:.4f}%".replace(".", ",")
                    problemas.append({"Base": nome, "Mês": df.loc[i, "data"].strftime("%m/%Y"),
                                      "Problema": f"a coluna de variação informa {fmt(v[i])}, mas o número-índice "
                                                  f"{n[i]:.9f} corresponde a {fmt(implicita[i])}. Confira qual dos "
                                                  "dois está errado (os cálculos usam o número-índice)."})
                    n_prob += 1
        resumo.append({"Base": nome, "Meses": len(df), "De": df["data"].min().strftime("%m/%Y"),
                       "Até": df["data"].max().strftime("%m/%Y"),
                       "Situação": "✅ consistente" if n_prob == 0 else f"❌ {n_prob} problema(s)"})
    return pd.DataFrame(resumo), pd.DataFrame(problemas, columns=["Base", "Mês", "Problema"])
